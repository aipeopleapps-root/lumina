from __future__ import annotations

import io
import logging

import openpyxl

from app.models import InputType

logger = logging.getLogger(__name__)

# Sheet name fingerprints that identify a Configuration Workbook
_CONFIG_SHEET_MARKERS = {
    "time off types",
    "time off plans",
    "leave of absence types",
    "holiday calendars",
    "period schedules",
    "absence type groups",
    "team absence (mss)",
    "team absence (ess)",
}

# Column header fingerprints for a User Stories Excel
_USER_STORY_COLS = {"as a", "i want", "so that"}


def _normalise(value: object) -> str:
    return str(value).strip().lower() if value is not None else ""


def detect_input_type(filename: str, file_bytes: bytes) -> InputType:
    """
    Infer the input type from the filename extension and workbook content.

    Detection order:
    1. .docx extension → bpw
    2. .xlsx with Config Workbook sheet names → config_workbook
    3. .xlsx with 'As a' / 'I want' / 'So that' columns → user_stories
    4. .xlsx with 'Requirement' column (or fallback) → requirements
    """
    lower_name = filename.lower()

    if lower_name.endswith(".docx"):
        logger.info("Detected input type: bpw (docx extension)")
        return InputType.bpw

    if lower_name.endswith(".xlsx"):
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            sheet_names_lower = {s.lower() for s in wb.sheetnames}
            wb.close()

            if sheet_names_lower & _CONFIG_SHEET_MARKERS:
                logger.info("Detected input type: config_workbook (matching sheet names)")
                return InputType.config_workbook

            # Re-open to inspect first sheet headers
            wb2 = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
            ws = wb2.active
            first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            wb2.close()

            headers = {_normalise(c) for c in first_row if c is not None}
            if _USER_STORY_COLS.issubset(headers):
                logger.info("Detected input type: user_stories (column headers)")
                return InputType.user_stories

        except Exception as exc:
            logger.warning("Could not inspect xlsx for type detection: %s", exc)

        logger.info("Detected input type: requirements (default xlsx)")
        return InputType.requirements

    # Unknown extension — attempt requirements as safe fallback
    logger.warning("Unrecognised file extension for '%s', defaulting to requirements", filename)
    return InputType.requirements
