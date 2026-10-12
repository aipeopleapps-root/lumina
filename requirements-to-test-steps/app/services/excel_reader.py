from __future__ import annotations

import io
import logging
from typing import Any

import openpyxl

from app.models import Requirement, UserStory

logger = logging.getLogger(__name__)

_COL_REQUIREMENT_ID = "requirement id"
_COL_REQUIREMENT = "requirement"
_COL_MODULE = "module"
_COL_PRIORITY = "priority"


def _normalise(value: Any) -> str:
    """Strip and lower-case a header cell value for comparison."""
    return str(value).strip().lower() if value is not None else ""


async def read_requirements(file_bytes: bytes) -> list[Requirement]:
    """Parse an xlsx file and return a list of Requirement objects.

    Raises ValueError if the 'Requirement' column is missing.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The uploaded Excel file is empty.")

    # Build column index map from header row
    header_row = rows[0]
    col_map: dict[str, int] = {}
    for idx, cell_value in enumerate(header_row):
        normalised = _normalise(cell_value)
        if normalised:
            col_map[normalised] = idx

    if _COL_REQUIREMENT not in col_map:
        raise ValueError(
            "Required column 'Requirement' not found in the uploaded file. "
            f"Found columns: {list(col_map.keys())}"
        )

    req_id_idx = col_map.get(_COL_REQUIREMENT_ID)
    req_text_idx = col_map[_COL_REQUIREMENT]
    module_idx = col_map.get(_COL_MODULE)
    priority_idx = col_map.get(_COL_PRIORITY)

    requirements: list[Requirement] = []

    for row_number, row in enumerate(rows[1:], start=2):
        # Skip blank rows (all cells are None or empty string)
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        req_text = row[req_text_idx] if req_text_idx < len(row) else None
        if req_text is None or str(req_text).strip() == "":
            logger.debug("Skipping row %d: empty Requirement cell", row_number)
            continue

        # Requirement ID
        raw_req_id = (
            row[req_id_idx]
            if (req_id_idx is not None and req_id_idx < len(row))
            else None
        )
        if raw_req_id is None or str(raw_req_id).strip() == "":
            req_id = f"REQ-{row_number:03d}"
        else:
            req_id = str(raw_req_id).strip()

        module = (
            str(row[module_idx]).strip()
            if (module_idx is not None and module_idx < len(row) and row[module_idx] is not None)
            else None
        )
        priority = (
            str(row[priority_idx]).strip()
            if (priority_idx is not None and priority_idx < len(row) and row[priority_idx] is not None)
            else None
        )

        requirements.append(
            Requirement(
                row_number=row_number,
                requirement_id=req_id,
                requirement_text=str(req_text).strip(),
                module=module if module else None,
                priority=priority if priority else None,
            )
        )

    logger.info("Parsed %d requirements from uploaded file", len(requirements))
    return requirements


async def read_user_stories(file_bytes: bytes) -> list[UserStory]:
    """
    Parse an xlsx User Stories file.

    Expected columns (case-insensitive): Story ID (optional), As a, I want,
    So that, Acceptance Criteria, Requirement ID (optional).
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The uploaded User Stories file is empty.")

    header_row = rows[0]
    col_map: dict[str, int] = {}
    for idx, cell_value in enumerate(header_row):
        normalised = _normalise(cell_value)
        if normalised:
            col_map[normalised] = idx

    # Flexible column matching
    def _find(patterns: list[str]) -> int | None:
        for p in patterns:
            if p in col_map:
                return col_map[p]
        for key, idx in col_map.items():
            for p in patterns:
                if p in key:
                    return idx
        return None

    as_a_idx = _find(["as a", "as_a"])
    i_want_idx = _find(["i want", "i_want"])
    so_that_idx = _find(["so that", "so_that"])
    story_id_idx = _find(["story id", "story_id"])
    req_id_idx = _find(["requirement id", "requirement_id"])
    ac_idx = _find(["acceptance criteria", "acceptance_criteria"])

    if as_a_idx is None or i_want_idx is None or so_that_idx is None:
        raise ValueError(
            "User Stories file must contain columns: 'As a', 'I want', 'So that'. "
            f"Found: {list(col_map.keys())}"
        )

    stories: list[UserStory] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if all(c is None or str(c).strip() == "" for c in row):
            continue

        def _get(idx: int | None) -> str:
            if idx is None or idx >= len(row) or row[idx] is None:
                return ""
            return str(row[idx]).strip()

        as_a = _get(as_a_idx)
        if not as_a:
            continue

        story_id = _get(story_id_idx) or f"US-{row_number - 1:03d}"
        req_id = _get(req_id_idx) or f"REQ-{row_number - 1:03d}"
        raw_ac = _get(ac_idx)
        ac_list = [c.strip() for c in raw_ac.replace("|", "\n").splitlines() if c.strip()]

        stories.append(UserStory(
            requirement_id=req_id,
            story_id=story_id,
            as_a=as_a,
            i_want=_get(i_want_idx),
            so_that=_get(so_that_idx),
            acceptance_criteria=ac_list,
        ))

    logger.info("Parsed %d user stories from uploaded file", len(stories))
    return stories
