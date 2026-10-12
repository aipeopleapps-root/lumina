from __future__ import annotations

import io
import logging

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import InputType, PipelineResult

logger = logging.getLogger(__name__)

_HEADER_FILL_HEX = "1F4E79"
_ALT_ROW_FILL_HEX = "D9E1F2"
_MAX_COL_WIDTH = 80


def _header_style() -> tuple[Font, PatternFill]:
    font = Font(bold=True, color="FFFFFF")
    fill = PatternFill(fill_type="solid", fgColor=_HEADER_FILL_HEX)
    return font, fill


def _alt_fill() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=_ALT_ROW_FILL_HEX)


def _autofit_columns(ws: object, max_width: int = _MAX_COL_WIDTH) -> None:
    """Set each column width to the max content length, capped at max_width."""
    col_widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                col_widths[cell.column] = max(
                    col_widths.get(cell.column, 0),
                    len(str(cell.value)),
                )
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, max_width)


def _write_sheet(
    ws: object,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    """Write headers + data rows with formatting to a worksheet."""
    header_font, header_fill = _header_style()
    alt = _alt_fill()

    # Write header row
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Write data rows
    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        if row_idx % 2 == 0:
            for cell in ws[row_idx]:
                cell.fill = alt
        for cell in ws[row_idx]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    _autofit_columns(ws)
    ws.freeze_panes = "A2"


def _build_sheet1_user_stories(wb: Workbook, result: PipelineResult) -> None:
    ws = wb.active
    ws.title = "User Stories"
    headers = ["Story ID", "Requirement ID", "As a", "I want", "So that", "Acceptance Criteria"]
    rows = [
        [us.story_id, us.requirement_id, us.as_a, us.i_want, us.so_that,
         " | ".join(us.acceptance_criteria)]
        for us in result.user_stories
    ]
    _write_sheet(ws, headers, rows)


def _build_sheet1_config_items(wb: Workbook, result: PipelineResult) -> None:
    ws = wb.active
    ws.title = "Config Items"
    headers = ["Item ID", "Workday Object", "Action", "Reference ID", "Display Name", "Key Configuration"]
    rows = [
        [
            item.item_id,
            item.sheet_name,
            item.action,
            item.reference_id,
            item.display_name,
            " | ".join(f"{k}: {v}" for k, v in item.additional_fields.items()),
        ]
        for item in result.config_items
    ]
    _write_sheet(ws, headers, rows)


def _build_sheet1_bpw_steps(wb: Workbook, result: PipelineResult) -> None:
    ws = wb.active
    ws.title = "Process Steps"
    headers = ["Process", "Step #", "Step Type", "Step Name", "Role", "Description", "Condition"]
    rows = [
        [
            step.process_name,
            step.step_number,
            step.step_type,
            step.step_name,
            step.role,
            step.description,
            step.condition or "",
        ]
        for step in result.bpw_steps
    ]
    _write_sheet(ws, headers, rows)


def build_output_workbook(result: PipelineResult) -> bytes:
    """Build a three-sheet Excel workbook and return it as bytes."""
    wb = Workbook()

    # ── Sheet 1: varies by input type ────────────────────────────────────────
    if result.input_type == InputType.config_workbook:
        _build_sheet1_config_items(wb, result)
    elif result.input_type == InputType.bpw:
        _build_sheet1_bpw_steps(wb, result)
    else:
        _build_sheet1_user_stories(wb, result)

    # ── Sheet 2: Test Scenarios ──────────────────────────────────────────────
    ws2 = wb.create_sheet("Test Scenarios")
    ref_label = {
        InputType.config_workbook: "Config Item",
        InputType.bpw: "Process",
    }.get(result.input_type, "Story ID")
    ts_headers = ["Scenario ID", ref_label, "Type", "Title", "Description", "Preconditions"]
    ts_rows = [
        [
            ts.scenario_id,
            ts.story_id,
            ts.scenario_type,
            ts.title,
            ts.description,
            " | ".join(ts.preconditions),
        ]
        for ts in result.test_scenarios
    ]
    _write_sheet(ws2, ts_headers, ts_rows)

    # ── Sheet 3: Test Steps ──────────────────────────────────────────────────
    ws3 = wb.create_sheet("Test Steps")
    step_headers = ["Step ID", "Scenario ID", "Step #", "Action", "Expected Result", "Test Data"]
    step_rows = [
        [
            step.step_id,
            step.scenario_id,
            step.step_number,
            step.action,
            step.expected_result,
            step.test_data or "",
        ]
        for step in result.test_steps
    ]
    _write_sheet(ws3, step_headers, step_rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    logger.info(
        "Built output workbook: %d user stories, %d scenarios, %d steps",
        len(result.user_stories),
        len(result.test_scenarios),
        len(result.test_steps),
    )
    return buffer.read()
