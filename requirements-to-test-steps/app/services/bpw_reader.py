from __future__ import annotations

import io
import logging
import re

from docx import Document
from docx.oxml.ns import qn

from app.models import BPWProcess, BPWStep

logger = logging.getLogger(__name__)

# Column positions in BP steps table (0-indexed)
_COL_STEP_NUM = 0
_COL_STEP_TYPE = 1
_COL_STEP_NAME = 2
_COL_ROLE = 3
_COL_DESC = 4
_COL_CONDITION = 5

_STEPS_TABLE_HEADERS = {"#", "step type", "step name", "role"}


def _is_steps_table(table: Any) -> bool:  # type: ignore[name-defined]
    """Identify a BP steps table by its first-row headers."""
    if not table.rows:
        return False
    first_row_text = {c.text.strip().lower() for c in table.rows[0].cells}
    return _STEPS_TABLE_HEADERS.issubset(first_row_text)


def _is_process_overview_table(table: Any) -> bool:  # type: ignore[name-defined]
    if not table.rows:
        return False
    headers = {c.text.strip().lower() for c in table.rows[0].cells}
    return "process" in headers and "process definition" in headers


def _is_approval_routes_table(table: Any) -> bool:  # type: ignore[name-defined]
    if not table.rows:
        return False
    first_row = {c.text.strip().lower() for c in table.rows[0].cells}
    return "route" in first_row and "initiator" in first_row


def _clean_step_num(raw: str) -> str:
    """Normalise step numbers: '77.' → '7.', strip extra whitespace."""
    cleaned = re.sub(r"(\d)\1+", r"\1", raw.strip())
    return cleaned or raw.strip()


def _extract_overview(table: Any) -> dict[str, str]:  # type: ignore[name-defined]
    overview: dict[str, str] = {}
    for row in table.rows[1:]:
        if len(row.cells) >= 2:
            key = row.cells[0].text.strip().lower()
            val = row.cells[1].text.strip()
            if key and val:
                overview[key] = val
    return overview


def _extract_approval_routes(table: Any) -> list[dict]:  # type: ignore[name-defined]
    routes = []
    for row in table.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if not any(cells):
            continue
        routes.append({
            "route": cells[0] if len(cells) > 0 else "",
            "initiator": cells[1] if len(cells) > 1 else "",
            "first_approval": cells[2] if len(cells) > 2 else "",
            "second_approval": cells[3] if len(cells) > 3 else "",
            "comments": cells[4] if len(cells) > 4 else "",
        })
    return routes


def _extract_steps(table: Any, process_name: str) -> list[BPWStep]:  # type: ignore[name-defined]
    steps = []
    for row in table.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if not any(cells):
            continue
        # Guard: skip rows that look like sub-headers (step_num non-numeric)
        step_num_raw = cells[_COL_STEP_NUM] if len(cells) > _COL_STEP_NUM else ""
        if not step_num_raw or not re.search(r"\d", step_num_raw):
            continue

        steps.append(BPWStep(
            process_name=process_name,
            step_number=_clean_step_num(step_num_raw),
            step_type=cells[_COL_STEP_TYPE] if len(cells) > _COL_STEP_TYPE else "",
            step_name=cells[_COL_STEP_NAME] if len(cells) > _COL_STEP_NAME else "",
            role=cells[_COL_ROLE] if len(cells) > _COL_ROLE else "",
            description=cells[_COL_DESC] if len(cells) > _COL_DESC else "",
            condition=cells[_COL_CONDITION] if len(cells) > _COL_CONDITION and cells[_COL_CONDITION] else None,
        ))
    return steps


def _body_children(doc: Any):  # type: ignore[name-defined]
    """Yield (element_type, object) tuples from document body in document order."""
    body = doc.element.body
    para_tag = qn("w:p")
    tbl_tag = qn("w:tbl")
    para_map = {p._element: p for p in doc.paragraphs}
    tbl_map = {t._element: t for t in doc.tables}

    for child in body:
        if child.tag == para_tag and child in para_map:
            yield "para", para_map[child]
        elif child.tag == tbl_tag and child in tbl_map:
            yield "table", tbl_map[child]


async def read_bpw(file_bytes: bytes) -> tuple[list[BPWStep], list[BPWProcess]]:
    """
    Parse a Business Process Workbook .docx.

    Returns:
        bpw_steps: flat list of all extracted steps (for UI review)
        bpw_processes: process metadata (overview, routes) for LLM context
    """
    doc = Document(io.BytesIO(file_bytes))

    overview_data: dict[str, str] = {}
    approval_routes: list[dict] = []
    # Map process_name → list of BPWStep
    process_steps: dict[str, list[BPWStep]] = {}
    last_heading: str = ""

    for elem_type, obj in _body_children(doc):
        if elem_type == "para":
            style = obj.style.name if obj.style else ""
            text = obj.text.strip()
            if not text:
                continue
            if "Heading" in style and text:
                # Strip "Process: " prefix if present
                last_heading = re.sub(r"^Process:\s*", "", text).strip()

        elif elem_type == "table":
            if _is_process_overview_table(obj):
                overview_data = _extract_overview(obj)
                logger.debug("Extracted process overview: %s", list(overview_data.keys()))

            elif _is_approval_routes_table(obj):
                approval_routes = _extract_approval_routes(obj)
                logger.debug("Extracted %d approval routes", len(approval_routes))

            elif _is_steps_table(obj):
                process_name = last_heading or "Unknown Process"
                steps = _extract_steps(obj, process_name)
                if process_name not in process_steps:
                    process_steps[process_name] = []
                process_steps[process_name].extend(steps)
                logger.info("Process '%s': extracted %d steps", process_name, len(steps))

    # Build BPWProcess objects — one per process, all sharing the same doc-level metadata
    bpw_processes: list[BPWProcess] = []
    for process_name, _steps in process_steps.items():
        bpw_processes.append(BPWProcess(
            process_name=process_name,
            overview=overview_data.get("overview"),
            objective=overview_data.get("objective"),
            business_trigger=overview_data.get("business trigger"),
            approval_routes=approval_routes,
        ))

    all_steps = [step for steps in process_steps.values() for step in steps]
    logger.info("BPW total: %d steps across %d processes", len(all_steps), len(bpw_processes))
    return all_steps, bpw_processes
