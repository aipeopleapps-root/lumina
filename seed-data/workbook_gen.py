"""
Generate a synthetic Workday Absence configuration workbook that looks like
a real Workday kit template. Includes realistic sheet names, column headers,
and example rules/logic that a consultant would recognize.
"""

from __future__ import annotations

import io
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill


def generate_absence_config_workbook(client_name: str) -> bytes:
    """
    Returns an in-memory .xlsx file (as bytes) for a fictional Absence
    configuration workbook for `client_name`.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    # Sheet 1: Absence Plans
    ws1 = wb.create_sheet("Absence Plans")
    _write_header(
        ws1,
        [
            "Plan Name",
            "Plan Code",
            "Unit of Time",
            "Accrual Frequency",
            "Carryover",
            "Max Carryover",
            "Effective Date",
            "Notes",
        ],
    )
    ws1.append(
        [
            "Annual Leave",
            "AL",
            "Days",
            "Monthly",
            "Yes",
            "5",
            "2024-01-01",
            "Standard annual leave plan",
        ]
    )
    ws1.append(
        [
            "Sick Leave",
            "SL",
            "Days",
            "Monthly",
            "No",
            "",
            "2024-01-01",
            "No carryover per policy",
        ]
    )
    ws1.append(
        [
            "Parental Leave",
            "PL",
            "Weeks",
            "Event-Based",
            "No",
            "",
            "2024-01-01",
            "Triggered by birth/adoption event",
        ]
    )

    # Sheet 2: Accrual Rules
    ws2 = wb.create_sheet("Accrual Rules")
    _write_header(
        ws2,
        [
            "Plan",
            "Eligibility Rule",
            "Service Band (Years)",
            "Accrual Rate",
            "Max Balance",
            "Proration",
            "Notes",
        ],
    )
    ws2.append(
        [
            "Annual Leave",
            "Active EE, >= 3 months service",
            "0-2",
            "1.25 days/month",
            "25",
            "Yes",
            "Prorated for new hires",
        ]
    )
    ws2.append(
        [
            "Annual Leave",
            "Active EE, >= 3 months service",
            "2-5",
            "1.67 days/month",
            "30",
            "Yes",
            "Higher rate after 2 years",
        ]
    )
    ws2.append(
        [
            "Annual Leave",
            "Active EE, >= 3 months service",
            "5+",
            "2.0 days/month",
            "35",
            "Yes",
            "Senior tenure rate",
        ]
    )
    ws2.append(
        [
            "Sick Leave",
            "All Active EE",
            "Any",
            "1.0 days/month",
            "12",
            "No",
            "Flat rate, no proration",
        ]
    )

    # Sheet 3: Eligibility Conditions (CRITICAL — often a source of anti-patterns)
    ws3 = wb.create_sheet("Eligibility Conditions")
    _write_header(
        ws3,
        [
            "Condition Name",
            "Plan",
            "Logic",
            "Applies To",
            "Priority",
            "Notes",
        ],
    )
    ws3.append(
        [
            "New Hire Waiting Period",
            "Annual Leave",
            "Hire Date + 90 days",
            "All new hires",
            "1",
            "Blocks accrual for first 3 months",
        ]
    )
    ws3.append(
        [
            "Manager Override Flag",
            "Annual Leave",
            "Custom field: Skip_Eligibility = Yes",
            "Exec / VP only",
            "2",
            "ANTI-PATTERN: Custom field instead of security group",
        ]
    )
    ws3.append(
        [
            "Contract Workers Excluded",
            "Sick Leave",
            "Employee Type != Contractor",
            "All EE",
            "1",
            "",
        ]
    )

    # Sheet 4: Time Off Business Process (where BP step orderings often confuse UAT)
    ws4 = wb.create_sheet("Time Off BP Steps")
    _write_header(
        ws4,
        [
            "Step",
            "Step Type",
            "Assigned To",
            "Condition",
            "Due Duration",
            "Notes",
        ],
    )
    ws4.append(
        [
            "1",
            "Initiate",
            "Employee",
            "Always",
            "",
            "Employee submits request",
        ]
    )
    ws4.append(
        [
            "2",
            "Approval",
            "Manager",
            "Request > 5 days OR Short Notice",
            "2 business days",
            "Manager approval required for long/urgent requests",
        ]
    )
    ws4.append(
        [
            "3",
            "Approval",
            "HR Partner",
            "Request > 10 days",
            "2 business days",
            "HR approval for extended leave",
        ]
    )
    ws4.append(
        [
            "4",
            "Notification",
            "Payroll",
            "Approved",
            "",
            "Auto-notify payroll on approval",
        ]
    )
    ws4.append(
        [
            "5",
            "Complete",
            "System",
            "Always",
            "",
            "Update balances",
        ]
    )

    # Sheet 5: Configuration Notes (where we put warnings / known issues)
    ws5 = wb.create_sheet("Config Notes")
    _write_header(
        ws5,
        [
            "Category",
            "Item",
            "Issue / Note",
            "Severity",
            "Owner",
        ],
    )
    ws5.append(
        [
            "Eligibility",
            "Manager Override Flag",
            "Using custom field for eligibility bypass is an anti-pattern. "
            "This will require manual updates and has no audit trail. "
            "Recommend replacing with a security group assignment.",
            "Medium",
            "Config Lead",
        ]
    )
    ws5.append(
        [
            "BP Design",
            "Step 2-3 Ordering",
            "Manager approval and HR approval are sequential. If manager is "
            "OOO, HR waits. Consider parallel routing for requests >10 days.",
            "Low",
            "BP Lead",
        ]
    )
    ws5.append(
        [
            "Carryover",
            "Max Carryover = 5 days",
            "Client policy states 5 days but legal in [Country] allows 10. "
            "Confirm with legal before go-live.",
            "High",
            "HR Policy Owner",
        ]
    )

    # Freeze panes, adjust column widths for readability.
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_hcm_config_workbook(client_name: str) -> bytes:
    """
    Returns an in-memory .xlsx for a fictional HCM Core configuration —
    org structures, job profiles, and compensation grades.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Sheet 1: Supervisory Orgs
    ws1 = wb.create_sheet("Supervisory Org Hierarchy")
    _write_header(
        ws1,
        [
            "Org Name",
            "Org Code",
            "Parent Org",
            "Manager",
            "Effective Date",
            "Inactive Date",
        ],
    )
    ws1.append([f"{client_name} Executive", "EXEC", "", "CEO Placeholder", "2024-01-01", ""])
    ws1.append(["Engineering", "ENG", f"{client_name} Executive", "VP Engineering", "2024-01-01", ""])
    ws1.append(["Product", "PROD", f"{client_name} Executive", "VP Product", "2024-01-01", ""])
    ws1.append(["Sales", "SALES", f"{client_name} Executive", "CRO Placeholder", "2024-01-01", ""])
    ws1.append(["Finance", "FIN", f"{client_name} Executive", "CFO Placeholder", "2024-01-01", ""])

    # Sheet 2: Job Profiles
    ws2 = wb.create_sheet("Job Profiles")
    _write_header(
        ws2,
        [
            "Job Profile",
            "Job Code",
            "Job Family",
            "Job Level",
            "FLSA Status",
            "Default Comp Grade",
        ],
    )
    ws2.append(["Software Engineer", "SWE", "Engineering", "IC3", "Exempt", "E3"])
    ws2.append(["Senior Software Engineer", "SWE-SR", "Engineering", "IC4", "Exempt", "E4"])
    ws2.append(["Engineering Manager", "EM", "Engineering", "M3", "Exempt", "M3"])
    ws2.append(["Product Manager", "PM", "Product", "IC3", "Exempt", "P3"])
    ws2.append(["Sales Rep", "REP", "Sales", "IC2", "Non-Exempt", "S2"])

    # Sheet 3: Compensation Grades
    ws3 = wb.create_sheet("Compensation Grades")
    _write_header(
        ws3,
        [
            "Grade",
            "Min Salary",
            "Mid Salary",
            "Max Salary",
            "Currency",
            "Effective Date",
            "Notes",
        ],
    )
    ws3.append(["E3", "80000", "100000", "120000", "USD", "2024-01-01", ""])
    ws3.append(["E4", "110000", "135000", "160000", "USD", "2024-01-01", ""])
    ws3.append(["M3", "130000", "155000", "180000", "USD", "2024-01-01", ""])
    ws3.append(["P3", "90000", "110000", "130000", "USD", "2024-01-01", ""])
    ws3.append(["S2", "50000", "65000", "80000", "USD", "2024-01-01", ""])

    # Sheet 4: Integration Mappings (common source of config issues)
    ws4 = wb.create_sheet("Integration Mappings")
    _write_header(
        ws4,
        [
            "Field Name",
            "Source System",
            "Source Field",
            "Workday Field",
            "Transformation",
            "Notes",
        ],
    )
    ws4.append(
        [
            "Employee ID",
            "Payroll System",
            "EMP_NUM",
            "Employee_ID",
            "Pad to 8 digits",
            "Legacy IDs are 6-digit, Workday requires 8",
        ]
    )
    ws4.append(
        [
            "Manager",
            "Payroll System",
            "MGR_ID",
            "Manager Reference",
            "Lookup by Employee_ID",
            "RISK: Circular manager refs will break load",
        ]
    )
    ws4.append(
        [
            "Hire Date",
            "HRIS",
            "START_DT",
            "Hire_Date",
            "Parse MM/DD/YYYY",
            "Source has inconsistent date formats",
        ]
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_header(ws, headers: list[str]):
    """Write a styled header row."""
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
