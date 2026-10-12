"""
Test Automation Suite — input detection, config-workbook reading, and the shared
execute path that both the Execute button and the nightly timer go through.

Workbooks are generated in-memory with openpyxl rather than committed as binaries,
matching test_pi_vsdx_parser.py's approach.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from docx import Document

from test_automation import handler, input_detector, xlsx_reader
from common.tenant import TenantContext


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _config_sheet_rows(module_specific_name: str) -> list[list]:
    return [
        ["Configuration Workbook"],
        [],
        ["Setup", "Action", "Reference ID", "Name", "Notes"],
        ["Setup", "Add", "REF-1", "Alpha", "first"],
        ["Setup", "Keep", "REF-2", "Beta", "skip me"],
        ["Setup", "Update", "REF-3", "Gamma", "second"],
        ["Setup", "Remove", "REF-4", "Delta", "third"],
    ]


def _docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Process: Onboard New Hire")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Input detection ──────────────────────────────────────────────────────────

def test_detects_config_workbook_for_a_module_not_in_any_allowlist():
    # Payroll-style sheet name — the old Absence-only allowlist returned "requirements" here.
    data = _xlsx({"Pay Groups": _config_sheet_rows("Pay Groups")})
    assert input_detector.detect_input_type("Payroll_Config.xlsx", data) == "config_workbook"


def test_detects_config_workbook_for_absence_sheets_too():
    data = _xlsx({"Time Off Types": _config_sheet_rows("Time Off Types")})
    assert input_detector.detect_input_type("Absence_Config.xlsx", data) == "config_workbook"


def test_config_detection_requires_a_testable_row():
    rows = [["Setup", "Action", "Reference ID", "Name"], ["Setup", "Keep", "REF-1", "Alpha"]]
    data = _xlsx({"Sheet1": rows})
    assert input_detector.detect_input_type("Only_Keep.xlsx", data) == "requirements"


def test_requirements_workbook_is_not_misdetected_as_config():
    rows = [
        ["Requirement ID", "Requirement", "Module", "Priority"],
        ["REQ-1", "Employee can request time off", "HCM", "High"],
    ]
    data = _xlsx({"Requirements": rows})
    assert input_detector.detect_input_type("Reqs.xlsx", data) == "requirements"


def test_detects_user_stories_workbook():
    rows = [["As a", "I want", "So that"], ["manager", "approve requests", "team stays staffed"]]
    data = _xlsx({"Stories": rows})
    assert input_detector.detect_input_type("Stories.xlsx", data) == "user_stories"


def test_detects_bpw_docx_by_extension():
    assert input_detector.detect_input_type("Process.docx", _docx()) == "bpw"


# ── Config reader ────────────────────────────────────────────────────────────

def test_read_config_workbook_keeps_add_update_remove_and_skips_keep():
    wb = openpyxl.load_workbook(io.BytesIO(_xlsx({"Pay Groups": _config_sheet_rows("Pay Groups")})))
    items = xlsx_reader.read_config_workbook(wb)
    assert [i["action"] for i in items] == ["Add", "Update", "Remove"]
    assert [i["reference_id"] for i in items] == ["REF-1", "REF-3", "REF-4"]
    assert all(i["sheet_name"] == "Pay Groups" for i in items)


# ── Shared execute path ──────────────────────────────────────────────────────

class _FakeQueue:
    def __init__(self):
        self.messages: list[dict] = []

    def enqueue_execution(self, payload: dict) -> None:
        self.messages.append(payload)


class _Audit:
    def __init__(self):
        self.events: list[tuple[str, dict]] = []

    def write(self, ctx, action, *, ip, resource_type, resource_id, resource_name, details):
        self.events.append((action, details))


@pytest.fixture
def ctx():
    return TenantContext(tenant_id="acme", user_id="user-1", role="consultant", email="c@example.com")


@pytest.fixture
def world(monkeypatch):
    """A ready run with two scenarios on an active environment, behind fake boundaries."""
    state = {
        "environment": {"environment_id": "env-1", "is_active": True, "tenant_id": "acme",
                        "nightly_enabled": True, "nightly_run_id": "run-1", "name": "SIT"},
        "run": {"run_id": "run-1", "status": "steps_ready", "file_name": "Reqs.xlsx",
                "execution_status": None},
        "credential": {"username": "svc"},
        "scenarios": [{"scenario_id": "sc-1"}, {"scenario_id": "sc-2"}],
        "status_updates": [],
    }
    queue = _FakeQueue()
    audit = _Audit()

    monkeypatch.setattr(handler.store, "get_environment", lambda **kw: state["environment"])
    monkeypatch.setattr(handler.store, "get_run", lambda **kw: state["run"])
    monkeypatch.setattr(handler.store, "list_scenarios", lambda **kw: state["scenarios"])
    monkeypatch.setattr(handler.store, "update_scenario_execution",
                        lambda **kw: state["status_updates"].append(kw))
    monkeypatch.setattr(handler.store, "update_run_execution", lambda **kw: state.setdefault("run_updates", []).append(kw))
    monkeypatch.setattr(handler.keyvault, "get_tenant_workday_credential",
                        lambda tenant_id, environment_id: state["credential"])
    monkeypatch.setattr(handler.queue, "enqueue_execution", queue.enqueue_execution)
    monkeypatch.setattr(handler.audit_events, "write", audit.write)
    return {"state": state, "queue": queue, "audit": audit}


def test_queue_run_execution_enqueues_every_scenario(world, ctx):
    result = handler.queue_run_execution(ctx=ctx, run_id="run-1", environment_id="env-1", ip="1.2.3.4", source="manual")
    assert result == {"error": None, "scenario_count": 2}
    assert [m["scenario_id"] for m in world["queue"].messages] == ["sc-1", "sc-2"]
    assert all(m["environment_id"] == "env-1" and m["tenant_id"] == "acme" for m in world["queue"].messages)
    assert world["audit"].events[0][0] == "test_automation.execute_triggered"
    assert world["audit"].events[0][1]["source"] == "manual"


@pytest.mark.parametrize("mutate, expected", [
    (lambda s: s["environment"].update(is_active=False), "environment not found"),
    (lambda s: s["run"].update(status="needs_review"), "no generated steps yet"),
    (lambda s: s["run"].update(execution_status="running"), "already executing"),
    (lambda s: s.update(credential=None), "no Workday credential"),
    (lambda s: s.update(scenarios=[]), "no scenarios"),
])
def test_queue_run_execution_refuses_unready_runs(world, ctx, mutate, expected):
    mutate(world["state"])
    result = handler.queue_run_execution(ctx=ctx, run_id="run-1", environment_id="env-1", ip="", source="manual")
    assert expected in result["error"]
    assert world["queue"].messages == []


# ── Nightly job ──────────────────────────────────────────────────────────────

def test_nightly_runs_every_scheduled_environment(world, monkeypatch):
    envs = [
        {"environment_id": "env-1", "tenant_id": "acme", "nightly_run_id": "run-1"},
        {"environment_id": "env-2", "tenant_id": "globex", "nightly_run_id": "run-9"},
    ]
    monkeypatch.setattr(handler.store, "list_nightly_environments_all_tenants", lambda: envs)
    # env-2 is a different tenant; the same fake run/env data serves both calls
    world["state"]["run"]["run_id"] = "run-1"
    monkeypatch.setattr(handler.store, "get_environment", lambda **kw: {"is_active": True})

    handler.run_nightly_executions()

    sources = {e[1]["source"] for e in world["audit"].events}
    assert sources == {"nightly"}
    assert len(world["audit"].events) == 2


def test_nightly_skips_an_environment_without_credentials_and_keeps_going(world, monkeypatch):
    envs = [
        {"environment_id": "env-1", "tenant_id": "acme", "nightly_run_id": "run-1"},
        {"environment_id": "env-2", "tenant_id": "acme", "nightly_run_id": "run-1"},
    ]
    monkeypatch.setattr(handler.store, "list_nightly_environments_all_tenants", lambda: envs)
    calls = {"n": 0}

    def cred(tenant_id, environment_id):
        calls["n"] += 1
        return None if environment_id == "env-1" else {"username": "svc"}

    monkeypatch.setattr(handler.keyvault, "get_tenant_workday_credential", cred)
    handler.run_nightly_executions()

    assert calls["n"] == 2
    assert len(world["audit"].events) == 1  # env-2 still ran


def test_nightly_does_nothing_when_no_environment_is_scheduled(world, monkeypatch):
    monkeypatch.setattr(handler.store, "list_nightly_environments_all_tenants", lambda: [])
    handler.run_nightly_executions()
    assert world["queue"].messages == []


# ── Internal routes: non-UUID ids are refused before any DB call ─────────────

def test_internal_scenario_route_rejects_non_uuid_id_without_touching_db(monkeypatch):
    # The execution worker's startup heartbeat calls GET .../scenarios/heartbeat-test.
    # Postgres would raise on a non-UUID; that surfaced as a 500 and made the health check lie.
    def boom(**kw):
        raise AssertionError("DB must not be called for a non-UUID id")
    monkeypatch.setattr(handler.store, "get_scenario", boom)
    event = {"pathParameters": {"scenario_id": "heartbeat-test"},
             "queryStringParameters": {"tenant_id": "heartbeat"}, "body": None}
    resp = handler._internal_get_scenario(event)
    assert resp["statusCode"] == 400


def test_internal_step_route_rejects_non_uuid_id(monkeypatch):
    monkeypatch.setattr(handler.store, "update_step_execution", lambda **kw: (_ for _ in ()).throw(AssertionError("no DB")))
    event = {"pathParameters": {"step_id": "not-a-uuid"}, "body": '{"tenant_id": "acme", "execution_status": "passed"}'}
    assert handler._internal_patch_step(event)["statusCode"] == 400


# ── Structured UI target on steps ────────────────────────────────────────────

from test_automation import agents
from common import test_automation_store as tas_store


def _scenario():
    return {"scenario_id": "sc-1", "scenario_type": "positive", "title": "Approve", "description": None,
            "preconditions": []}


def test_generate_test_steps_keeps_ui_target(monkeypatch):
    monkeypatch.setattr(agents.ai, "generate_json", lambda *a, **k: {"steps": [
        {"step_number": 1, "action": "Click the 'Create New' button", "ui_target": "Create New",
         "expected_result": "form opens", "test_data": None},
        {"step_number": 2, "action": "Log in as the Finance Manager", "ui_target": None,
         "expected_result": "dashboard shows", "test_data": None},
    ]})
    steps = agents.generate_test_steps(_scenario())
    assert steps[0]["ui_target"] == "Create New"
    assert steps[1]["ui_target"] is None


def test_generate_test_steps_treats_blank_ui_target_as_none(monkeypatch):
    monkeypatch.setattr(agents.ai, "generate_json", lambda *a, **k: {"steps": [
        {"step_number": 1, "action": "Open the page", "ui_target": "  ",
         "expected_result": "page opens", "test_data": None},
    ]})
    assert agents.generate_test_steps(_scenario())[0]["ui_target"] is None


def test_step_row_maps_ui_target_and_tolerates_old_rows_without_it():
    row = ("step-1", "sc-1", 1, "Click the 'Save' button", "saved", None, "not_executed", None, None, "Save")
    assert tas_store._step_row_to_dict(row)["ui_target"] == "Save"
    old_row = row[:9] + (None,)
    assert tas_store._step_row_to_dict(old_row)["ui_target"] is None


def test_ui_target_column_is_added_once_per_process():
    calls = []

    class Cur:
        def execute(self, sql, *a):
            calls.append(sql)

    tas_store._ui_target_column_ready = False
    tas_store._ensure_ui_target_column(Cur())
    tas_store._ensure_ui_target_column(Cur())
    tas_store._ui_target_column_ready = False
    assert len(calls) == 1 and "ADD COLUMN IF NOT EXISTS ui_target" in calls[0]
