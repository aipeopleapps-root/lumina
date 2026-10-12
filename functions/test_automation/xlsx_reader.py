"""
Ported from requirements-to-test-steps/app/services/excel_reader.py +
config_workbook_reader.py, combined into one module since both are plain
openpyxl readers over an uploaded .xlsx. Sync (no asyncio), returns plain
dicts instead of Pydantic models — matches the rest of Lumina's Python
handlers, none of which use pydantic.
"""

from __future__ import annotations

import logging
from typing import Any

import openpyxl

log = logging.getLogger(__name__)

_COL_REQUIREMENT_ID = "requirement id"
_COL_REQUIREMENT = "requirement"
_COL_MODULE = "module"
_COL_PRIORITY = "priority"


def _normalise(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def read_requirements(wb: openpyxl.Workbook) -> list[dict]:
    """Parse the active sheet as a Requirements workbook. Raises ValueError
    if the 'Requirement' column is missing."""
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The uploaded Excel file is empty.")

    header_row = rows[0]
    col_map: dict[str, int] = {}
    for idx, cell_value in enumerate(header_row):
        normalised = _normalise(cell_value)
        if normalised:
            col_map[normalised] = idx

    if _COL_REQUIREMENT not in col_map:
        raise ValueError(
            "Required column 'Requirement' not found in the uploaded file. "
            f"Found columns: {list(col_map.keys())}"
        )

    req_id_idx = col_map.get(_COL_REQUIREMENT_ID)
    req_text_idx = col_map[_COL_REQUIREMENT]
    module_idx = col_map.get(_COL_MODULE)
    priority_idx = col_map.get(_COL_PRIORITY)

    requirements: list[dict] = []

    for row_number, row in enumerate(rows[1:], start=2):
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue

        req_text = row[req_text_idx] if req_text_idx < len(row) else None
        if req_text is None or str(req_text).strip() == "":
            continue

        raw_req_id = (
            row[req_id_idx] if (req_id_idx is not None and req_id_idx < len(row)) else None
        )
        req_id = str(raw_req_id).strip() if raw_req_id and str(raw_req_id).strip() else f"REQ-{row_number:03d}"

        module = (
            str(row[module_idx]).strip()
            if (module_idx is not None and module_idx < len(row) and row[module_idx] is not None)
            else None
        )
        priority = (
            str(row[priority_idx]).strip()
            if (priority_idx is not None and priority_idx < len(row) and row[priority_idx] is not None)
            else None
        )

        requirements.append({
            "row_number": row_number,
            "requirement_id": req_id,
            "requirement_text": str(req_text).strip(),
            "module": module or None,
            "priority": priority or None,
        })

    log.info("Parsed %d requirements from uploaded file", len(requirements))
    return requirements


def read_user_stories(wb: openpyxl.Workbook) -> list[dict]:
    """
    Parse the active sheet as a User Stories workbook. Expected columns
    (case-insensitive, flexibly matched): Story ID (optional), As a, I want,
    So that, Acceptance Criteria, Requirement ID (optional).
    """
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The uploaded User Stories file is empty.")

    header_row = rows[0]
    col_map: dict[str, int] = {}
    for idx, cell_value in enumerate(header_row):
        normalised = _normalise(cell_value)
        if normalised:
            col_map[normalised] = idx

    def _find(patterns: list[str]) -> int | None:
        for p in patterns:
            if p in col_map:
                return col_map[p]
        for key, idx in col_map.items():
            for p in patterns:
                if p in key:
                    return idx
        return None

    as_a_idx = _find(["as a", "as_a"])
    i_want_idx = _find(["i want", "i_want"])
    so_that_idx = _find(["so that", "so_that"])
    story_id_idx = _find(["story id", "story_id"])
    req_id_idx = _find(["requirement id", "requirement_id"])
    ac_idx = _find(["acceptance criteria", "acceptance_criteria"])

    if as_a_idx is None or i_want_idx is None or so_that_idx is None:
        raise ValueError(
            "User Stories file must contain columns: 'As a', 'I want', 'So that'. "
            f"Found: {list(col_map.keys())}"
        )

    stories: list[dict] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if all(c is None or str(c).strip() == "" for c in row):
            continue

        def _get(idx: int | None) -> str:
            if idx is None or idx >= len(row) or row[idx] is None:
                return ""
            return str(row[idx]).strip()

        as_a = _get(as_a_idx)
        if not as_a:
            continue

        story_id = _get(story_id_idx) or f"US-{row_number - 1:03d}"
        req_id = _get(req_id_idx) or f"REQ-{row_number - 1:03d}"
        raw_ac = _get(ac_idx)
        ac_list = [c.strip() for c in raw_ac.replace("|", "\n").splitlines() if c.strip()]

        stories.append({
            "requirement_id": req_id,
            "story_id": story_id,
            "as_a": as_a,
            "i_want": _get(i_want_idx),
            "so_that": _get(so_that_idx),
            "acceptance_criteria": ac_list,
        })

    log.info("Parsed %d user stories from uploaded file", len(stories))
    return stories


# ── Config workbook ─────────────────────────────────────────────────────────

_SKIP_SHEETS = {"coversheet", "instructions & contents", "staging >>>", "output >>>", "sheet20"}
_TESTABLE_ACTIONS = {"add", "update", "remove"}


def _norm(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def _find_header_row(ws: Any) -> tuple[int, int, list[str]] | None:
    """
    Scan rows 1-25 for the actual column header row. Criteria: contains
    'setup' or 'action' in the first five columns AND has >=3 non-empty
    cells, with the immediately adjacent column also non-empty (rules out
    sparse section-header rows).
    """
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=25, values_only=True)):
        non_empty = sum(1 for c in row if c is not None and str(c).strip())
        if non_empty < 3:
            continue
        headers = [str(v).strip() if v is not None else "" for v in row]
        candidates = {_norm(row[c]): c for c in range(min(5, len(row))) if _norm(row[c]) in ("setup", "action")}
        # "Action" holds the Add/Update/Remove values; "Setup" can sit to its left as a
        # label column, so prefer Action whenever both are present.
        for name in ("action", "setup"):
            col_idx = candidates.get(name)
            if col_idx is None:
                continue
            next_col = col_idx + 1
            if next_col >= len(row) or row[next_col] is None or not str(row[next_col]).strip():
                continue
            return row_idx, col_idx, headers
    return None


def _pick_display_name(row: tuple, headers: list[str], action_col: int, ref_col: int) -> str:
    priority_cols = [
        i for i, h in enumerate(headers)
        if i not in (action_col, ref_col) and i < len(row)
        and ("name" in h.lower() or "type" in h.lower())
        and row[i] is not None
        and str(row[i]).strip()
    ]
    if priority_cols:
        return str(row[priority_cols[0]]).strip()

    next_col = ref_col + 1
    if next_col < len(row) and row[next_col] is not None:
        return str(row[next_col]).strip()
    return ""


def _find_ref_col(headers: list[str]) -> int | None:
    for i, h in enumerate(headers):
        if h.lower().startswith("reference id"):
            return i
    return None


def _parse_sheet(ws: Any, sheet_name: str, counter: list[int]) -> list[dict]:
    result = _find_header_row(ws)
    if result is None:
        return []

    header_row_idx, action_col, headers = result
    ref_col = _find_ref_col(headers)

    items: list[dict] = []
    for row in ws.iter_rows(min_row=header_row_idx + 2, values_only=True):
        if all(c is None or str(c).strip() == "" for c in row):
            continue

        action_raw = row[action_col] if action_col < len(row) else None
        if action_raw is None:
            continue
        action = str(action_raw).strip()
        if action.lower() not in _TESTABLE_ACTIONS:
            continue

        display_name = _pick_display_name(row, headers, action_col, ref_col or 0)

        ref_id = ""
        if ref_col is not None and ref_col < len(row) and row[ref_col] is not None:
            ref_id = str(row[ref_col]).strip()
        if not ref_id:
            ref_id = display_name.upper().replace(" ", "_") if display_name else f"UNKNOWN_ROW_{counter[0]}"

        additional: dict[str, str] = {}
        for col_idx, (header, cell) in enumerate(zip(headers, row)):
            if col_idx in (action_col, ref_col):
                continue
            if header and cell is not None and str(cell).strip():
                val = str(cell).strip()
                if len(val) <= 200:
                    additional[header] = val

        item_id = f"CFG-{counter[0]:03d}"
        counter[0] += 1

        items.append({
            "item_id": item_id,
            "sheet_name": sheet_name,
            "action": action,
            "reference_id": ref_id,
            "display_name": display_name or ref_id,
            "additional_fields": additional,
        })

    return items


def has_config_structure(wb: openpyxl.Workbook) -> bool:
    """True if any non-skipped sheet has a Setup/Action header row AND at least one
    testable (Add/Update/Remove) row — structural, so it holds for any module's
    configuration workbook, not only Absence's sheet names."""
    for sheet_name in wb.sheetnames:
        if sheet_name.lower() in _SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        result = _find_header_row(ws)
        if result is None:
            continue
        header_row_idx, action_col, _ = result
        for row in ws.iter_rows(min_row=header_row_idx + 2, values_only=True):
            if action_col < len(row) and row[action_col] is not None:
                if str(row[action_col]).strip().lower() in _TESTABLE_ACTIONS:
                    return True
    return False


def read_config_workbook(wb: openpyxl.Workbook) -> list[dict]:
    """
    Parse a Configuration Workbook (any module) and return its testable config items.
    Only Add, Update, and Remove rows are included (Keep and Sample skipped).
    """
    counter = [1]
    all_items: list[dict] = []

    for sheet_name in wb.sheetnames:
        if sheet_name.lower() in _SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        items = _parse_sheet(ws, sheet_name, counter)
        if items:
            log.info("Sheet '%s': found %d testable config items", sheet_name, len(items))
        all_items.extend(items)

    log.info("Config workbook total: %d testable items across %d sheets", len(all_items), len(wb.sheetnames))
    return all_items
