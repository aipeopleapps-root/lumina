"""
Ported near-verbatim from requirements-to-test-steps/app/services/bpw_reader.py
(Business Process Workbook .docx table-structure walking). Sync, plain dicts.
"""

from __future__ import annotations

import logging
import re
from typing import Any

from docx import Document
from docx.oxml.ns import qn

log = logging.getLogger(__name__)

_COL_STEP_NUM = 0
_COL_STEP_TYPE = 1
_COL_STEP_NAME = 2
_COL_ROLE = 3
_COL_DESC = 4
_COL_CONDITION = 5

_STEPS_TABLE_HEADERS = {"#", "step type", "step name", "role"}

# Second recognized steps-table shape: the BP Workbook Generation app's own
# Section 7 output (bp-workbook-api/src/lib/docxBuilder.js) uses different
# column names/order ("Step Order"/"Step"/"Type"/.../"Security Group" instead
# of "#"/"Step Name"/"Step Type"/"Role") since the two were built
# independently — see docs/AZURE_MIGRATION.md's 2026-10-01 entry. Detected
# and mapped separately rather than forcing one shared header convention.
_BPW_GEN_STEPS_HEADERS = {"step order", "step", "type", "description"}
_BPW_GEN_COL_STEP_NUM = 0     # "Step Order" — a letter (A, B, C…), not numeric
_BPW_GEN_COL_STEP_NAME = 1    # "Step"
_BPW_GEN_COL_STEP_TYPE = 2    # "Type"
_BPW_GEN_COL_DESC = 3         # "Description"
_BPW_GEN_COL_CONDITION = 4    # "Condition Rules"
_BPW_GEN_COL_ROLE = 5         # "Security Group" — closest analog to a Role/responsible-party column


def _is_steps_table(table: Any) -> bool:
    if not table.rows:
        return False
    first_row_text = {c.text.strip().lower() for c in table.rows[0].cells}
    return _STEPS_TABLE_HEADERS.issubset(first_row_text)


def _is_bpw_gen_steps_table(table: Any) -> bool:
    if not table.rows:
        return False
    first_row_text = {c.text.strip().lower() for c in table.rows[0].cells}
    return _BPW_GEN_STEPS_HEADERS.issubset(first_row_text)


def _is_process_overview_table(table: Any) -> bool:
    if not table.rows:
        return False
    headers = {c.text.strip().lower() for c in table.rows[0].cells}
    return "process" in headers and "process definition" in headers


def _is_approval_routes_table(table: Any) -> bool:
    if not table.rows:
        return False
    first_row = {c.text.strip().lower() for c in table.rows[0].cells}
    return "route" in first_row and "initiator" in first_row


def _clean_step_num(raw: str) -> str:
    cleaned = re.sub(r"(\d)\1+", r"\1", raw.strip())
    return cleaned or raw.strip()


def _extract_overview(table: Any) -> dict[str, str]:
    overview: dict[str, str] = {}
    for row in table.rows[1:]:
        if len(row.cells) >= 2:
            key = row.cells[0].text.strip().lower()
            val = row.cells[1].text.strip()
            if key and val:
                overview[key] = val
    return overview


def _extract_approval_routes(table: Any) -> list[dict]:
    routes = []
    for row in table.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if not any(cells):
            continue
        routes.append({
            "route": cells[0] if len(cells) > 0 else "",
            "initiator": cells[1] if len(cells) > 1 else "",
            "first_approval": cells[2] if len(cells) > 2 else "",
            "second_approval": cells[3] if len(cells) > 3 else "",
            "comments": cells[4] if len(cells) > 4 else "",
        })
    return routes


def _extract_steps(table: Any, process_name: str) -> list[dict]:
    steps = []
    for row in table.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if not any(cells):
            continue
        step_num_raw = cells[_COL_STEP_NUM] if len(cells) > _COL_STEP_NUM else ""
        if not step_num_raw or not re.search(r"\d", step_num_raw):
            continue

        steps.append({
            "process_name": process_name,
            "step_number": _clean_step_num(step_num_raw),
            "step_type": cells[_COL_STEP_TYPE] if len(cells) > _COL_STEP_TYPE else "",
            "step_name": cells[_COL_STEP_NAME] if len(cells) > _COL_STEP_NAME else "",
            "role": cells[_COL_ROLE] if len(cells) > _COL_ROLE else "",
            "description": cells[_COL_DESC] if len(cells) > _COL_DESC else "",
            "condition": cells[_COL_CONDITION] if len(cells) > _COL_CONDITION and cells[_COL_CONDITION] else None,
        })
    return steps


def _extract_bpw_gen_steps(table: Any, process_name: str) -> list[dict]:
    steps = []
    for row in table.rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if not any(cells):
            continue
        step_name = cells[_BPW_GEN_COL_STEP_NAME] if len(cells) > _BPW_GEN_COL_STEP_NAME else ""
        if not step_name:
            # Unlike the native TAS shape, "Step Order" here is a letter
            # (A, B, C…) with no digit to validate against — an empty Step
            # name is the real signal a row is a blank placeholder, not data.
            continue
        condition = cells[_BPW_GEN_COL_CONDITION] if len(cells) > _BPW_GEN_COL_CONDITION else ""
        steps.append({
            "process_name": process_name,
            "step_number": cells[_BPW_GEN_COL_STEP_NUM] if len(cells) > _BPW_GEN_COL_STEP_NUM else "",
            "step_type": cells[_BPW_GEN_COL_STEP_TYPE] if len(cells) > _BPW_GEN_COL_STEP_TYPE else "",
            "step_name": step_name,
            "role": cells[_BPW_GEN_COL_ROLE] if len(cells) > _BPW_GEN_COL_ROLE else "",
            "description": cells[_BPW_GEN_COL_DESC] if len(cells) > _BPW_GEN_COL_DESC else "",
            "condition": condition if condition and condition.upper() != "N/A" else None,
        })
    return steps


def _body_children(doc: Any):
    body = doc.element.body
    para_tag = qn("w:p")
    tbl_tag = qn("w:tbl")
    para_map = {p._element: p for p in doc.paragraphs}
    tbl_map = {t._element: t for t in doc.tables}

    for child in body:
        if child.tag == para_tag and child in para_map:
            yield "para", para_map[child]
        elif child.tag == tbl_tag and child in tbl_map:
            yield "table", tbl_map[child]


def read_bpw(file_bytes: bytes) -> tuple[list[dict], list[dict]]:
    """
    Parse a Business Process Workbook .docx.

    Returns:
        bpw_steps: flat list of all extracted steps (for UI review)
        bpw_processes: process metadata (overview, routes) for LLM context
    """
    import io
    doc = Document(io.BytesIO(file_bytes))

    overview_data: dict[str, str] = {}
    approval_routes: list[dict] = []
    process_steps: dict[str, list[dict]] = {}
    last_heading: str = ""

    for elem_type, obj in _body_children(doc):
        if elem_type == "para":
            style = obj.style.name if obj.style else ""
            text = obj.text.strip()
            if not text:
                continue
            if "Heading" in style and text:
                last_heading = re.sub(r"^Process:\s*", "", text).strip()

        elif elem_type == "table":
            if _is_process_overview_table(obj):
                overview_data = _extract_overview(obj)

            elif _is_approval_routes_table(obj):
                approval_routes = _extract_approval_routes(obj)

            elif _is_steps_table(obj):
                process_name = last_heading or overview_data.get("process") or "Unknown Process"
                steps = _extract_steps(obj, process_name)
                process_steps.setdefault(process_name, []).extend(steps)
                log.info("Process '%s': extracted %d steps", process_name, len(steps))

            elif _is_bpw_gen_steps_table(obj):
                process_name = last_heading or overview_data.get("process") or "Unknown Process"
                steps = _extract_bpw_gen_steps(obj, process_name)
                process_steps.setdefault(process_name, []).extend(steps)
                log.info("Process '%s': extracted %d steps (BP Workbook Generation shape)", process_name, len(steps))

    bpw_processes: list[dict] = []
    for process_name in process_steps:
        bpw_processes.append({
            "process_name": process_name,
            "overview": overview_data.get("overview"),
            "objective": overview_data.get("objective"),
            "business_trigger": overview_data.get("business trigger"),
            "approval_routes": approval_routes,
        })

    all_steps = [step for steps in process_steps.values() for step in steps]
    log.info("BPW total: %d steps across %d processes", len(all_steps), len(bpw_processes))
    return all_steps, bpw_processes
