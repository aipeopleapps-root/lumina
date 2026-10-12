"""
Ported from requirements-to-test-steps/app/services/excel_writer.py — same
3-sheet formatting (dark-blue header, alternating row shading, frozen header,
autofit). Adapted to read from plain dicts sourced from Postgres
(test_automation_store.list_scenarios/list_steps + runs.input_items) instead
of a Pydantic PipelineResult, and to take pre-computed display labels
("US-001"/"TS-001"/"TS-001-001") rather than trusting model-generated IDs.
"""

from __future__ import annotations

import io
import logging

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

_HEADER_FILL_HEX = "1F4E79"
_ALT_ROW_FILL_HEX = "D9E1F2"
_MAX_COL_WIDTH = 80


def _header_style() -> tuple[Font, PatternFill]:
    return Font(bold=True, color="FFFFFF"), PatternFill(fill_type="solid", fgColor=_HEADER_FILL_HEX)


def _alt_fill() -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=_ALT_ROW_FILL_HEX)


def _autofit_columns(ws, max_width: int = _MAX_COL_WIDTH) -> None:
    col_widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                col_widths[cell.column] = max(col_widths.get(cell.column, 0), len(str(cell.value)))
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, max_width)


def _write_sheet(ws, headers: list[str], rows: list[list[str]]) -> None:
    header_font, header_fill = _header_style()
    alt = _alt_fill()

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        if row_idx % 2 == 0:
            for cell in ws[row_idx]:
                cell.fill = alt
        for cell in ws[row_idx]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    _autofit_columns(ws)
    ws.freeze_panes = "A2"


def build_output_workbook(
    *,
    input_type: str,
    sheet1_items: list[dict],
    scenarios: list[dict],
    steps: list[dict],
) -> bytes:
    """
    Build a three-sheet Excel workbook and return it as bytes.

    sheet1_items shape varies by input_type:
      requirements/user_stories -> {story_label, requirement_id, as_a, i_want, so_that, acceptance_criteria}
      config_workbook -> {item_id, sheet_name, action, reference_id, display_name, additional_fields}
      bpw -> {process_name, step_number, step_type, step_name, role, description, condition}
    scenarios: {scenario_label, source_ref, scenario_type, title, description, preconditions}
    steps: {step_label, scenario_label, step_number, action, expected_result, test_data}
    """
    wb = Workbook()

    if input_type == "config_workbook":
        ws1 = wb.active
        ws1.title = "Config Items"
        headers = ["Item ID", "Workday Object", "Action", "Reference ID", "Display Name", "Key Configuration"]
        rows = [
            [
                item["item_id"], item["sheet_name"], item["action"], item["reference_id"],
                item["display_name"],
                " | ".join(f"{k}: {v}" for k, v in (item.get("additional_fields") or {}).items()),
            ]
            for item in sheet1_items
        ]
        ref_label = "Config Item"
    elif input_type == "bpw":
        ws1 = wb.active
        ws1.title = "Process Steps"
        headers = ["Process", "Step #", "Step Type", "Step Name", "Role", "Description", "Condition"]
        rows = [
            [
                step["process_name"], step["step_number"], step["step_type"], step["step_name"],
                step["role"], step["description"], step.get("condition") or "",
            ]
            for step in sheet1_items
        ]
        ref_label = "Process"
    else:
        ws1 = wb.active
        ws1.title = "User Stories"
        headers = ["Story ID", "Requirement ID", "As a", "I want", "So that", "Acceptance Criteria"]
        rows = [
            [
                us["story_label"], us["requirement_id"], us["as_a"], us["i_want"], us["so_that"],
                " | ".join(us.get("acceptance_criteria") or []),
            ]
            for us in sheet1_items
        ]
        ref_label = "Story ID"
    _write_sheet(ws1, headers, rows)

    ws2 = wb.create_sheet("Test Scenarios")
    ts_headers = ["Scenario ID", ref_label, "Type", "Title", "Description", "Preconditions"]
    ts_rows = [
        [
            sc["scenario_label"], sc["source_ref"], sc["scenario_type"], sc["title"],
            sc["description"], " | ".join(sc.get("preconditions") or []),
        ]
        for sc in scenarios
    ]
    _write_sheet(ws2, ts_headers, ts_rows)

    ws3 = wb.create_sheet("Test Steps")
    step_headers = ["Step ID", "Scenario ID", "Step #", "Action", "Expected Result", "Test Data"]
    step_rows = [
        [
            st["step_label"], st["scenario_label"], st["step_number"], st["action"],
            st["expected_result"], st.get("test_data") or "",
        ]
        for st in steps
    ]
    _write_sheet(ws3, step_headers, step_rows)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    log.info(
        "Built output workbook: %d sheet1 items, %d scenarios, %d steps",
        len(sheet1_items), len(scenarios), len(steps),
    )
    return buffer.read()


def build_execution_report_workbook(
    *,
    run: dict,
    scenarios: list[dict],
    steps: list[dict],
) -> bytes:
    """
    Build a three-sheet Workday-execution results workbook: Summary, Scenario
    Results, Step Results. Separate from build_output_workbook() above (which
    exports the generated test *script*, not execution outcomes) — see
    handler.py's _export_results.

    scenarios: {scenario_label, title, scenario_type, execution_status,
                execution_started_at, execution_completed_at}
    steps: {step_label, scenario_label, action, expected_result,
            execution_status, execution_result}
    """
    wb = Workbook()

    terminal = {"passed", "failed", "error"}
    counts = {"passed": 0, "failed": 0, "error": 0, "pending": 0, "running": 0, "not_executed": 0}
    for sc in scenarios:
        counts[sc.get("execution_status") or "not_executed"] = (
            counts.get(sc.get("execution_status") or "not_executed", 0) + 1
        )
    executed = sum(v for k, v in counts.items() if k in terminal)
    pass_rate = f"{(counts['passed'] / executed * 100):.1f}%" if executed else "n/a"

    ws1 = wb.active
    ws1.title = "Summary"
    _write_sheet(
        ws1,
        ["Field", "Value"],
        [
            ["Run", run.get("file_name") or run.get("run_id")],
            ["Environment", run.get("environment_id") or "n/a"],
            ["Executed", run.get("updated_at") or "n/a"],
            ["Total scenarios", str(len(scenarios))],
            ["Passed", str(counts["passed"])],
            ["Failed", str(counts["failed"])],
            ["Error", str(counts["error"])],
            ["Pending / running", str(counts["pending"] + counts["running"])],
            ["Pass rate (of executed)", pass_rate],
        ],
    )

    ws2 = wb.create_sheet("Scenario Results")
    _write_sheet(
        ws2,
        ["Scenario ID", "Title", "Type", "Status", "Started", "Completed"],
        [
            [
                sc["scenario_label"], sc["title"], sc["scenario_type"],
                sc.get("execution_status") or "not_executed",
                sc.get("execution_started_at") or "", sc.get("execution_completed_at") or "",
            ]
            for sc in scenarios
        ],
    )

    ws3 = wb.create_sheet("Step Results")
    _write_sheet(
        ws3,
        ["Step ID", "Scenario ID", "Action", "Expected Result", "Status", "Note"],
        [
            [
                st["step_label"], st["scenario_label"], st["action"], st["expected_result"],
                st.get("execution_status") or "not_executed",
                (st.get("execution_result") or {}).get("note") or "",
            ]
            for st in steps
        ],
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    log.info(
        "Built execution report workbook: %d scenarios, %d steps, pass_rate=%s",
        len(scenarios), len(steps), pass_rate,
    )
    return buffer.read()
